#!/usr/bin/env python3
"""
fill_leads.py — 用 curl 查 Lead API + 填表单提交 + 中文名 + Excel。

用法：
  cd "/Users/seeed/Documents/landing page traffic analysis"
  python3 scripts/fill_leads.py
"""
import csv
import datetime as dt
import json
import re
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
KEY_FILE = Path("/Users/seeed/Desktop/solution落地页/GA4/LeadApi/searchKey.txt")
HOST = "https://rel-search.seeedstudio.com"
CST = dt.timezone(dt.timedelta(hours=8))

SLUG_CN = {
    "smart-warehouse-management": "智能仓储管理",
    "voicecollectionanalysis": "语音采集分析",
    "conversational-voice-ai": "对话式语音AI",
    "smart-agriculture-sensing": "智慧农业感知",
    "smart-livestock-farming": "智慧畜牧",
    "intelligent-video-analytics": "智能视频分析",
    "indoor-outdoor-positioning": "室内外定位",
    "environment-monitoring": "环境监测",
    "building-energy-management": "楼宇能源管理",
    "campus-safety-management": "校园安全管理",
    "building-energy-retrofit": "楼宇节能改造",
    "hazard-response": "应急响应",
    "off-grid-mesh-communication": "离网通信",
    "acoustic-audio-ai": "声学与音频AI",
    "home-building-automation": "家庭与楼宇自动化",
}

def week_ending(date_str):
    d = dt.date.fromisoformat(date_str)
    return (d + dt.timedelta(days=(6 - d.weekday()) % 7)).isoformat()

def load_key():
    with open(KEY_FILE, encoding="utf-8") as f:
        return f.read().strip()

def curl_search(key, params):
    """用 curl 调 Typesense API，绕过 Python SSL 问题。"""
    body = json.dumps({"searches": [params]})
    result = subprocess.run(
        ["curl", "-s", "-X", "POST", f"{HOST}/multi_search",
         "-H", "Content-Type: application/json",
         "-H", f"X-TYPESENSE-API-KEY: {key}",
         "-d", body],
        capture_output=True, text=True, timeout=60
    )
    if result.returncode != 0:
        print(f"  curl error: {result.stderr[:200]}")
        return None
    try:
        payload = json.loads(result.stdout)
    except json.JSONDecodeError:
        print(f"  JSON parse error: {result.stdout[:200]}")
        return None
    r = payload["results"][0]
    if r.get("error"):
        print(f"  API error: {r['error']}")
        return None
    return r

def search_leads(key):
    """导出全部 leads。"""
    all_docs = []
    upper = int(dt.datetime.now(tz=CST).timestamp() * 1000)
    page = 1
    while True:
        params = {
            "collection": "leads", "q": "*", "query_by": "reason_for_contact",
            "per_page": 250, "page": page,
            "sort_by": "create_at_timestamp:desc",
            "filter_by": f"create_at_timestamp:<={upper}",
        }
        r = curl_search(key, params)
        if not r:
            break
        hits = r.get("hits", [])
        all_docs.extend(h["document"] for h in hits)
        total = r["found"]
        print(f"  page {page}: got {len(hits)}, total {len(all_docs)}/{total}", flush=True)
        if not hits or len(all_docs) >= total:
            break
        page += 1
    return all_docs

def match_slug(referrer, pageLocation=""):
    for slug, cn in SLUG_CN.items():
        pattern = f"/solutions/{re.escape(slug)}(-zh-hans)?(?![a-z])"
        for url in (referrer or "", pageLocation or ""):
            if re.search(pattern, url, re.IGNORECASE):
                return cn
    return None

def count_leads(docs):
    counts = defaultdict(int)
    total_by_week = defaultdict(int)
    matched = 0
    for doc in docs:
        referrer = doc.get("referrer", "")
        pageLocation = doc.get("pageLocation", "")
        create_at = doc.get("create_at", "")
        create_date = create_at[:10] if create_at else ""
        if not create_date:
            continue
        we = week_ending(create_date)
        total_by_week[we] += 1
        cn = match_slug(referrer, pageLocation)
        if cn:
            counts[(we, cn)] += 1
            matched += 1
    has_pl = sum(1 for d in docs if d.get("pageLocation"))
    sol_pl = sum(1 for d in docs if "/solutions/" in d.get("pageLocation", ""))
    print(f"\nLeads total: {len(docs)}, matched to solutions: {matched}")
    print(f"  pageLocation: {has_pl} have it, {sol_pl} contain /solutions/")
    if sol_pl:
        print("  pageLocation samples with /solutions/:")
        shown = 0
        for d in docs:
            pl = d.get("pageLocation", "")
            if "/solutions/" in pl and shown < 5:
                print(f"    {pl[:120]}")
                shown += 1
    return counts, total_by_week

def update_tables(leads_by_week, total_leads_by_week):
    analysis = ROOT / "data" / "analysis"
    
    # 表1
    t1_path = analysis / "方案增长总览.csv"
    with open(t1_path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        t1_fields = list(reader.fieldnames) + ["form_submits"]
        t1_rows = list(reader)
    for r in t1_rows:
        r["form_submits"] = total_leads_by_week.get(r["week_ending"], 0)
    with open(t1_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=t1_fields)
        w.writeheader(); w.writerows(t1_rows)
    print(f"Updated: {t1_path.name}")

    # 表2
    t2_path = analysis / "单方案流量明细.csv"
    with open(t2_path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        t2_fields = reader.fieldnames
        t2_rows = list(reader)
    for r in t2_rows:
        cn = SLUG_CN.get(r["slug"], r["slug"])
        r["slug"] = cn
        r["form_submits"] = leads_by_week.get((r["week_ending"], cn), 0)
        r["data_status"] = "GA4已接入·Lead已接入"
    with open(t2_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=t2_fields)
        w.writeheader(); w.writerows(t2_rows)
    print(f"Updated: {t2_path.name}")

    # 表3
    t3_path = analysis / "转换漏斗.csv"
    with open(t3_path, newline="", encoding="utf-8") as f:
        reader = csv.DictReader(f)
        t3_fields = reader.fieldnames
        t3_rows = list(reader)
    for r in t3_rows:
        cn = SLUG_CN.get(r["slug"], r["slug"])
        r["slug"] = cn
        r["form_submits"] = leads_by_week.get((r["week_ending"], cn), 0)
        r["data_status"] = "GA4部分·Lead已接入·CRM未接入"
    with open(t3_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=t3_fields)
        w.writeheader(); w.writerows(t3_rows)
    print(f"Updated: {t3_path.name}")

def generate_excel():
    import shutil
    OUT = ROOT / "exports/feishu-import"
    OUT.mkdir(parents=True, exist_ok=True)
    
    COL_MAPS = {
        "方案增长总览.csv": {
            "week_ending":"周次","solution_users":"方案页独立访客",
            "solution_pv":"Solution页面总访问量","cta_clicks":"CTA总点击量",
            "cta_click_rate":"CTA点击率","key_events":"关键事件",
            "key_event_conv_rate":"关键事件转化率",
            "fastest_growing_solution":"增长最快的solution",
            "top_traffic_solution":"流量最高方案",
            "social_impressions":"社媒总曝光量","solution_wow_pct":"Solution页面周环比",
            "form_submits":"表单提交总数",
            "lead_rate":"留资转化率",
        },
        "单方案流量明细.csv": {
            "week_ending":"周次","slug":"方案名称","landing_pv":"落地访问量",
            "users":"独立访客数","sessions":"session","avg_eng_s":"平均停留时长(秒)",
            "engagement_rate":"参与率","top_channel":"主要流量来源",
            "cta_clicks":"CTA点击量","form_submits":"表单提交",
            "key_events":"event","data_status":"数据状态","wow_pct":"周环比(%)",
        },
        "转换漏斗.csv": {
            "week_ending":"最新周次","slug":"方案名称","page_pv":"页面访问量",
            "cta_clicks":"CTA点击量","form_submits":"表单提交","add_to_cart":"一键加购",
            "pv_to_cta_rate":"页面到CTA转化率","cta_to_form_rate":"CTA到表单转化率",
            "data_status":"数据状态","engaged_sessions":"参与会话数",
        },
    }
    
    for fname, mapping in COL_MAPS.items():
        src = ROOT / "data/analysis" / fname
        rows_out = []
        with open(src, newline="", encoding="utf-8") as f:
            reader = csv.DictReader(f)
            new_fields = [mapping.get(h, h) for h in reader.fieldnames]
            for r in reader:
                rows_out.append({mapping.get(k, k): v for k, v in r.items()})
        with open(OUT / fname, "w", newline="", encoding="utf-8-sig") as f:
            w = csv.DictWriter(f, fieldnames=new_fields)
            w.writeheader(); w.writerows(rows_out)
        print(f"CSV: {fname} ({len(rows_out)} rows)")

    try:
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
        from openpyxl.utils import get_column_letter
        
        hf = Font(name="Microsoft YaHei", bold=True, size=11, color="FFFFFF")
        hfill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
        ha = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cf = Font(name="Microsoft YaHei", size=10)
        ca = Alignment(horizontal="left", vertical="center")
        tb = Border(left=Side(style="thin",color="D9D9D9"),right=Side(style="thin",color="D9D9D9"),
                     top=Side(style="thin",color="D9D9D9"),bottom=Side(style="thin",color="D9D9D9"))
        
        wb = Workbook()
        wb.remove(wb.active)
        for fname, sheet_name in [("方案增长总览.csv","方案增长总览"),("单方案流量明细.csv","单方案流量明细"),("转换漏斗.csv","转换漏斗")]:
            with open(OUT / fname, newline="", encoding="utf-8-sig") as f:
                rows = list(csv.reader(f))
            ws = wb.create_sheet(title=sheet_name)
            headers, data = rows[0], rows[1:]
            for ci, h in enumerate(headers, 1):
                c = ws.cell(row=1, column=ci, value=h)
                c.font = hf; c.fill = hfill; c.alignment = ha; c.border = tb
            for ri, row in enumerate(data, 2):
                for ci, val in enumerate(row, 1):
                    try:
                        num = float(val) if "." in val else (int(val) if val.lstrip("-").isdigit() else val)
                    except (ValueError, AttributeError):
                        num = val
                    c = ws.cell(row=ri, column=ci, value=num if num != "" else None)
                    c.font = cf; c.alignment = ca; c.border = tb
            for ci in range(1, len(headers)+1):
                ml = len(str(headers[ci-1]))
                for row in data:
                    if ci-1 < len(row): ml = max(ml, len(str(row[ci-1])))
                ws.column_dimensions[get_column_letter(ci)].width = min(ml*1.8+2, 30)
            ws.freeze_panes = "A2"
            print(f"Sheet '{sheet_name}': {len(data)} rows")
        
        xlsx = ROOT / "exports/feishu-import/GA4周度报表_三表合一.xlsx"
        wb.save(xlsx)
        shutil.copy(xlsx, "/tmp/ga4_weekly_tables.xlsx")
        print(f"\nExcel: {xlsx} ({xlsx.stat().st_size} bytes)")
        print("Copied to /tmp/ga4_weekly_tables.xlsx")
    except ImportError:
        print("openpyxl not available — 用 bundled Python 重跑或 pip install openpyxl")

def main():
    print("=== 1. 查询 Lead API (via curl) ===")
    key = load_key()
    docs = search_leads(key)
    if not docs:
        print("ERROR: 未获取到 leads 数据")
        return
    
    print("\n=== 2. 统计 leads 按方案×周 ===")
    leads_by_week, total_leads_by_week = count_leads(docs)
    
    print("\n=== 3. 更新三张表 ===")
    update_tables(leads_by_week, total_leads_by_week)
    
    print("\n=== 4. 生成 Excel ===")
    generate_excel()
    
    print("\n=== 完成 ===")
    print("文件:")
    print(f"  data/analysis/*.csv")
    print(f"  exports/feishu-import/*.csv")
    print(f"  exports/feishu-import/GA4周度报表_三表合一.xlsx")
    print(f"  /tmp/ga4_weekly_tables.xlsx")

if __name__ == "__main__":
    main()

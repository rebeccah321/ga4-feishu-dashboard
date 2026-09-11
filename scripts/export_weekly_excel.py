#!/usr/bin/env python3
"""
export_weekly_excel.py — 从 weekly_tables.json 导出 W23~W36 三张周度表 Excel。

口径与 generate_tables_correct.py 一致：
  - 每个 solution 单独查 GA4：pagePath CONTAINS /solutions/<slug>
  - 维度 date / hostName / sessionDefaultChannelGroup
  - 三语言站合并，周结束日为周日
  - CRM 表单提交与 lead_rate 由 merge_leads.py 回填

输出：
  exports/Seeed周度三表_W23-W36_<latest>.xlsx
  三张 sheet：方案增长总览 / 单方案流量明细 / 转换漏斗
"""
from __future__ import annotations

import datetime as dt
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "dashboard" / "data" / "weekly" / "weekly_tables.json"
EXPORTS = ROOT / "exports"

HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def iso_week(value: str) -> str:
    try:
        return f"W{dt.date.fromisoformat(str(value)).isocalendar()[1]}"
    except ValueError:
        return str(value)


def load() -> dict:
    if not SRC.exists():
        raise SystemExit(f"missing {SRC}")
    return json.loads(SRC.read_text(encoding="utf-8"))


def style_sheet(ws, headers: list[str], rows: list[list]):
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for row in rows:
        ws.append(row)
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
        for cell in row:
            cell.border = BORDER
    for idx, col in enumerate(ws.columns, start=1):
        max_len = max((len(str(c.value)) if c.value is not None else 0) for c in col)
        ws.column_dimensions[get_column_letter(idx)].width = min(max(max_len + 2, 10), 38)
    ws.freeze_panes = "A2"


def overview_rows(payload: dict) -> list[list]:
    out = []
    for row in payload.get("overview", []):
        we = row["week_ending"]
        out.append([
            iso_week(we),
            we,
            row.get("solution_users"),
            row.get("solution_pv"),
            row.get("cta_clicks"),
            row.get("cta_click_rate"),
            row.get("solution_wow_pct"),
            row.get("top_traffic_solution"),
            row.get("fastest_growing_solution"),
        ])
    return out


def detail_rows(payload: dict) -> list[list]:
    out = []
    for row in payload.get("detail", []):
        out.append([
            iso_week(row["week_ending"]),
            row["week_ending"],
            row.get("solution_name"),
            row.get("landing_pv"),
            row.get("users"),
            row.get("sessions"),
            row.get("avg_eng_s"),
            row.get("engagement_rate_pct"),
            row.get("top_channel"),
            row.get("cta_clicks"),
            row.get("form_submits"),
            row.get("wow_pct"),
            row.get("data_status"),
        ])
    return out


def funnel_rows(payload: dict) -> list[list]:
    out = []
    for row in payload.get("funnel", []):
        page_pv = row.get("page_pv") or 0
        form_submits = row.get("form_submits") or 0
        form_rate = round(100 * form_submits / page_pv, 2) if page_pv else 0
        out.append([
            iso_week(row["week_ending"]),
            row["week_ending"],
            row.get("solution_name"),
            row.get("page_pv"),
            row.get("sessions"),
            row.get("cta_clicks"),
            row.get("form_submits"),
            row.get("pv_to_cta_rate"),
            form_rate,
            row.get("data_status"),
        ])
    return out


def main() -> None:
    payload = load()
    latest = payload.get("latest_week_ending", "unknown")
    out_path = EXPORTS / f"Seeed周度三表_W23-W36_{latest}.xlsx"
    EXPORTS.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws1 = wb.active
    ws1.title = "方案增长总览"
    style_sheet(ws1, [
        "周次", "截止日期", "方案页独立访客数", "Solution页面总访问量", "CTA点击量",
        "CTA点击率(%)", "Solution页面周环比", "流量最高方案", "增长最快方案",
    ], overview_rows(payload))

    ws2 = wb.create_sheet("单方案流量明细")
    style_sheet(ws2, [
        "周次", "截止日期", "方案名称", "落地页访问量", "独立访客数",
        "Session", "平均停留时长（秒）", "参与率", "主要流量来源",
        "CTA点击量", "leads", "周环比", "数据状态",
    ], detail_rows(payload))

    ws3 = wb.create_sheet("转换漏斗")
    style_sheet(ws3, [
        "周次", "截止日期", "方案名称", "页面访问量（PV）", "Session",
        "CTA点击量", "表单提交（Leads）", "CTA点击率(%)", "表单转化率(%)",
        "数据状态",
    ], funnel_rows(payload))

    wb.save(out_path)
    print(f"wrote {out_path}")
    print(f"rows: overview={len(payload.get('overview', []))}, "
          f"detail={len(payload.get('detail', []))}, funnel={len(payload.get('funnel', []))}")


if __name__ == "__main__":
    main()
